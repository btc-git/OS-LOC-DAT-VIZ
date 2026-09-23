"""
Open Source Location Data Visualizer - github.com/btc-git/OS-LOC-DAT-VIZ
Licensed under the GNU General Public License v3.0 - see LICENSE file for details
"""

import os
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import Workbook
from PyQt6.QtCore import QDateTime, QTimer
from PyQt6.QtTest import QSignalSpy
from PyQt6.QtWidgets import QApplication

from import_wizard import (
    FIXED_UTC_OFFSETS,
    ImportWizardDialog,
    NAMED_TIMEZONE_CHOICES,
    SourceFileLoadWorker,
)
from kml_generator import KMLGenerator


class ImportWizardDialogUITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        cls.app = QApplication.instance() or QApplication([])

    def test_timezone_defaults_and_target_no_change_option(self):
        dialog = ImportWizardDialog()
        source_items = [dialog.source_timezone_combo.itemText(i) for i in range(dialog.source_timezone_combo.count())]
        target_items = [dialog.target_timezone_combo.itemText(i) for i in range(dialog.target_timezone_combo.count())]

        self.assertEqual("Fixed UTC+00:00", dialog.source_timezone_combo.currentText())
        self.assertEqual("No Change", dialog.target_timezone_combo.currentText())
        self.assertEqual((None, None), dialog.timezone_selection(dialog.target_timezone_combo))
        self.assertEqual(source_items, target_items[1:])
        self.assertEqual(len(NAMED_TIMEZONE_CHOICES) + len(FIXED_UTC_OFFSETS), len(source_items))
        self.assertNotIn("UTC (UTC+00:00)", source_items)
        self.assertIn("Fixed UTC+00:00", source_items)
        self.assertGreater(
            source_items.index("Fixed UTC+00:00"),
            source_items.index("Hawaii (UTC-10:00)"),
        )
        self.assertIn("US Eastern (UTC-05:00 / UTC-04:00 DST)", source_items)
        self.assertIn("Fixed UTC-05:00", source_items)
        self.assertIn("Fixed UTC+05:30", source_items)

    def test_new_source_resets_timezone_defaults(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "records.csv"
            pd.DataFrame([{
                "Timestamp": "2024-01-15 14:00:00",
                "Latitude": 43.15,
                "Longitude": -77.61,
            }]).to_csv(path, index=False)
            dialog = ImportWizardDialog()
            dialog.source_timezone_combo.setCurrentIndex(1)
            dialog.target_timezone_combo.setCurrentIndex(2)

            dialog.load_file(str(path), background=False)

            self.assertEqual(
                "Fixed UTC+00:00", dialog.source_timezone_combo.currentText()
            )
            self.assertEqual("No Change", dialog.target_timezone_combo.currentText())

    def test_timezone_choices_ignore_mouse_wheel(self):
        dialog = ImportWizardDialog()

        for combo in (dialog.source_timezone_combo, dialog.target_timezone_combo):
            starting_index = combo.currentIndex()
            wheel_event = Mock()

            combo.wheelEvent(wheel_event)

            self.assertEqual(starting_index, combo.currentIndex())
            wheel_event.ignore.assert_called_once_with()

    def test_source_loader_returns_excel_sheets_and_raw_records(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "records.xlsx"
            workbook = Workbook()
            workbook.active.title = "Records"
            workbook.active.append(["Timestamp", "Latitude", "Longitude"])
            workbook.active.append(["2024-01-15 14:00:00", 43.15, -77.61])
            workbook.create_sheet("Other")
            workbook.save(path)

            result = SourceFileLoadWorker.read_source(path)

            self.assertEqual(["Records", "Other"], result['sheet_names'])
            self.assertEqual("Records", result['sheet_name'])
            self.assertEqual(
                ["Timestamp", "Latitude", "Longitude"],
                result['raw_dataframe'].iloc[0].tolist(),
            )

    def test_dialog_loads_source_in_background_with_busy_progress(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "records.csv"
            pd.DataFrame([{
                "Timestamp": "2024-01-15 14:00:00",
                "Latitude": 43.15,
                "Longitude": -77.61,
            }]).to_csv(path, index=False)
            dialog = ImportWizardDialog()
            loading_finished = QSignalSpy(dialog.source_loading_finished)
            reader_started = threading.Event()
            release_reader = threading.Event()
            original_read_source = SourceFileLoadWorker.read_source

            def delayed_read_source(*args, **kwargs):
                reader_started.set()
                release_reader.wait(5)
                return original_read_source(*args, **kwargs)

            with patch.object(
                SourceFileLoadWorker,
                "read_source",
                side_effect=delayed_read_source,
            ):
                dialog.load_file(str(path))
                while not reader_started.is_set():
                    self.app.processEvents()

                heartbeat = []
                QTimer.singleShot(0, lambda: heartbeat.append(True))
                self.app.processEvents()

                self.assertEqual([True], heartbeat)
                self.assertFalse(dialog.source_load_progress.isHidden())
                self.assertFalse(dialog.header_row_spinbox.isEnabled())
                release_reader.set()
                self.assertTrue(loading_finished.wait(5000))
            self.app.processEvents()
            self.assertEqual(1, len(dialog.source_dataframe))
            self.assertTrue(dialog.source_load_progress.isHidden())
            self.assertTrue(dialog.header_row_spinbox.isEnabled())

    def test_numeric_date_order_and_filter_display_formats(self):
        dialog = ImportWizardDialog()
        labels = [label.text() for label in dialog.findChildren(type(dialog.mapping_note))]

        self.assertIn("Date format in source records:", labels)
        self.assertIn(
            "Choose the date order shown in the mapped source timestamp. "
            "A likely order is selected automatically; confirm it against the source records.",
            labels,
        )
        self.assertEqual(
            "Year-Month-Day (YYYY-MM-DD or YYYY/MM/DD)",
            dialog.date_order_combo.itemText(0),
        )
        self.assertEqual(
            "Year-Day-Month (YYYY-DD-MM or YYYY/DD/MM)",
            dialog.date_order_combo.itemText(1),
        )
        self.assertEqual(['YMD', 'YDM', 'MDY', 'DMY'], [
            dialog.date_order_combo.itemData(index)
            for index in range(dialog.date_order_combo.count())
        ])
        self.assertEqual("MM/dd/yyyy", dialog.filter_start_edit.displayFormat())
        self.assertEqual("MM/dd/yyyy", dialog.filter_end_edit.displayFormat())

        dialog.date_order_combo.setCurrentIndex(
            dialog.date_order_combo.findData('YMD')
        )
        self.assertEqual("yyyy-MM-dd", dialog.filter_start_edit.displayFormat())

        dialog.filter_enabled_checkbox.setChecked(True)
        dialog.filter_exact_times_checkbox.setChecked(True)
        self.assertEqual("yyyy-MM-dd h:mm:ss AP", dialog.filter_start_edit.displayFormat())
        self.assertEqual("yyyy-MM-dd h:mm:ss AP", dialog.filter_end_edit.displayFormat())

    def test_filter_uses_display_timezone_or_source_when_unchanged(self):
        dialog = ImportWizardDialog()
        dialog.filter_enabled_checkbox.setChecked(True)
        dialog.date_order_combo.setCurrentIndex(
            dialog.date_order_combo.findData('YMD')
        )
        selected_day = QDateTime.fromString(
            '2024-01-15 12:00:00', 'yyyy-MM-dd HH:mm:ss'
        )
        records = pd.DataFrame([
            {
                'Timestamp': '2024-01-15T02:00:00Z',
                'Latitude': 43.15,
                'Longitude': -77.61,
            },
            {
                'Timestamp': '2024-01-16T02:00:00Z',
                'Latitude': 43.16,
                'Longitude': -77.62,
            },
        ])

        eastern_index = next(
            index
            for index in range(dialog.target_timezone_combo.count())
            if (
                dialog.target_timezone_combo.itemData(index) or {}
            ).get('timezone_name') == 'America/New_York'
        )
        dialog.target_timezone_combo.setCurrentIndex(eastern_index)
        dialog.filter_start_edit.setDateTime(selected_day)
        dialog.filter_end_edit.setDateTime(selected_day)

        display_filtered, display_metadata = dialog.apply_datetime_filter(
            records
        )

        self.assertEqual([1], display_filtered.index.tolist())
        self.assertEqual('display', display_metadata['filter_timezone_basis'])
        self.assertEqual('2024-01-15T05:00:00Z', display_metadata['start_utc'])
        self.assertEqual('2024-01-16T04:59:59Z', display_metadata['end_utc'])
        self.assertIn(
            'selected display timezone', dialog.filter_timezone_note.text()
        )

        dialog.target_timezone_combo.setCurrentIndex(0)
        dialog.filter_start_edit.setDateTime(selected_day)
        dialog.filter_end_edit.setDateTime(selected_day)

        source_filtered, source_metadata = dialog.apply_datetime_filter(records)

        self.assertEqual([0], source_filtered.index.tolist())
        self.assertEqual('source', source_metadata['filter_timezone_basis'])
        self.assertEqual('2024-01-15T00:00:00Z', source_metadata['start_utc'])
        self.assertEqual('2024-01-15T23:59:59Z', source_metadata['end_utc'])
        self.assertIn(
            'source timezone because display is No Change',
            dialog.filter_timezone_note.text(),
        )

    def test_coordinate_labels_distinguish_tower_and_location_records(self):
        dialog = ImportWizardDialog()

        for data_type in ("Tower/Sector", "Distance from Tower"):
            dialog.record_type_combo.setCurrentIndex(
                dialog.record_type_combo.findData(data_type)
            )
            self.assertEqual(
                "Cell tower/site latitude",
                dialog.mapping_labels["Latitude"].text(),
            )
            self.assertEqual(
                "Cell tower/site longitude",
                dialog.mapping_labels["Longitude"].text(),
            )

        dialog.record_type_combo.setCurrentIndex(
            dialog.record_type_combo.findData("Location Point")
        )
        self.assertEqual("Latitude", dialog.mapping_labels["Latitude"].text())
        self.assertEqual("Longitude", dialog.mapping_labels["Longitude"].text())

    def test_excel_header_row_8_preview_uses_1_based_physical_rows(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "header_row_8.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            for row_number in range(1, 8):
                sheet.append([f"metadata {row_number}"])
            sheet.append(["Timestamp", "Latitude", "Longitude", "Azimuth"])
            for idx in range(32):
                sheet.append([
                    f"2024-01-15 14:{idx:02d}:00",
                    43.15 + idx * 0.001,
                    -77.61 - idx * 0.001,
                    180,
                ])
            workbook.save(path)

            with patch.object(
                SourceFileLoadWorker,
                "read_source",
                wraps=SourceFileLoadWorker.read_source,
            ) as read_source:
                dialog = ImportWizardDialog()
                dialog._header_reload_timer.setInterval(0)
                dialog.load_file(str(path), background=False)
                with patch.object(
                    dialog, "populate_preview", wraps=dialog.populate_preview
                ) as populate_preview:
                    for header_row in range(2, 9):
                        dialog.header_row_spinbox.setValue(header_row)
                    self.app.processEvents()

            self.assertEqual(["Timestamp", "Latitude", "Longitude", "Azimuth"], list(dialog.source_dataframe.columns))
            self.assertEqual(0, dialog.source_dataframe.index[0])
            self.assertIn("Header row: 8", dialog.preview_label.text())
            self.assertEqual(25, dialog.preview_table.rowCount())
            self.assertEqual("9", dialog.preview_table.verticalHeaderItem(0).text())
            self.assertEqual(1, read_source.call_count)
            self.assertEqual(1, populate_preview.call_count)

    def test_mapped_timestamps_set_date_order_and_filter_endpoints(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "timestamp_range.csv"
            pd.DataFrame([
                {"Timestamp": "2024-01-15 14:00:00", "Latitude": 43.15, "Longitude": -77.61},
                {"Timestamp": "2024-01-16 15:15:00", "Latitude": 43.16, "Longitude": -77.62},
                {"Timestamp": "2024-01-17 16:30:00", "Latitude": 43.17, "Longitude": -77.63},
            ]).to_csv(path, index=False)

            dialog = ImportWizardDialog()
            dialog.load_file(str(path), background=False)

            self.assertEqual('YMD', dialog.date_order_combo.currentData())
            self.assertEqual("yyyy-MM-dd", dialog.filter_start_edit.displayFormat())
            self.assertEqual(
                "2024-01-15 14:00:00",
                dialog.filter_start_edit.dateTime().toString("yyyy-MM-dd HH:mm:ss"),
            )
            self.assertEqual(
                "2024-01-17 16:30:00",
                dialog.filter_end_edit.dateTime().toString("yyyy-MM-dd HH:mm:ss"),
            )

            eastern_index = next(
                index
                for index in range(dialog.target_timezone_combo.count())
                if (
                    dialog.target_timezone_combo.itemData(index) or {}
                ).get('timezone_name') == 'America/New_York'
            )
            dialog.target_timezone_combo.setCurrentIndex(eastern_index)

            self.assertEqual(
                "2024-01-15 09:00:00",
                dialog.filter_start_edit.dateTime().toString(
                    "yyyy-MM-dd HH:mm:ss"
                ),
            )
            self.assertEqual(
                "2024-01-17 11:30:00",
                dialog.filter_end_edit.dateTime().toString(
                    "yyyy-MM-dd HH:mm:ss"
                ),
            )

    def test_timestamp_defaults_do_not_parse_every_source_row(self):
        dialog = ImportWizardDialog()
        row_count = 5000
        dialog.source_dataframe = pd.DataFrame({
            "Timestamp": ["2024-01-15 14:00:00"] * row_count,
            "Latitude": [43.15] * row_count,
            "Longitude": [-77.61] * row_count,
        })
        dialog.populate_mapping_options()

        parse_call_count = 0
        original_parse = KMLGenerator.parse_timestamp_to_kml

        def counting_parse(generator, timestamp):
            nonlocal parse_call_count
            parse_call_count += 1
            return original_parse(generator, timestamp)

        with patch.object(
            KMLGenerator, "parse_timestamp_to_kml", new=counting_parse
        ):
            dialog.populate_filter_range_from_source()

        self.assertEqual(2, parse_call_count)

    def test_year_day_month_filter_uses_detected_source_order(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "year_day_month.csv"
            pd.DataFrame([
                {"Timestamp": "2024-15-01 14:00:00", "Latitude": 43.15, "Longitude": -77.61},
                {"Timestamp": "2024-17-01 16:30:00", "Latitude": 43.17, "Longitude": -77.63},
            ]).to_csv(path, index=False)

            dialog = ImportWizardDialog()
            dialog.load_file(str(path), background=False)
            dialog.filter_enabled_checkbox.setChecked(True)
            dialog.filter_exact_times_checkbox.setChecked(True)
            filtered, metadata = dialog.apply_datetime_filter(
                dialog.build_normalized_dataframe(dialog.current_mappings())
            )

            self.assertEqual('YDM', dialog.date_order_combo.currentData())
            self.assertEqual("yyyy-dd-MM h:mm:ss AP", dialog.filter_start_edit.displayFormat())
            self.assertEqual(2, len(filtered))
            self.assertEqual(2, metadata['retained_rows'])

    def test_csv_preview_row_numbers_are_1_based_physical_rows(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "records.csv"
            pd.DataFrame([
                {"Timestamp": "2024-01-15 14:00:00", "Latitude": 43.15, "Longitude": -77.61},
                {"Timestamp": "2024-01-15 14:15:00", "Latitude": 43.16, "Longitude": -77.62},
            ]).to_csv(path, index=False)

            dialog = ImportWizardDialog()
            dialog.load_file(str(path), background=False)
            dialog.header_row_spinbox.setValue(1)

            self.assertIn("Header row: 1", dialog.preview_label.text())
            self.assertEqual("2", dialog.preview_table.verticalHeaderItem(0).text())

    def test_cell_site_lookup_uses_composite_key_and_field_fallback(self):
        records = pd.DataFrame([
            {
                "Start_eNodeB": 1001,
                "Start_Sector": 1,
                "Latitude": 40.1,
                "Longitude": -70.1,
                "Azimuth": 10,
            },
            {
                "Start_eNodeB": 1001,
                "Start_Sector": 2,
                "Latitude": 40.2,
                "Longitude": -70.2,
                "Azimuth": 20,
            },
            {
                "Start_eNodeB": 9999,
                "Start_Sector": 1,
                "Latitude": 40.3,
                "Longitude": -70.3,
                "Azimuth": 30,
            },
            {
                "Start_eNodeB": None,
                "Start_Sector": 1,
                "Latitude": 40.4,
                "Longitude": -70.4,
                "Azimuth": 40,
            },
        ])
        cell_sites = pd.DataFrame([
            {
                "E/G NodeB ID": 1001,
                "Cell ID": 1,
                "Site Lat": 45.1,
                "Site Long": -75.1,
                "Azimuth": 120,
            },
            {
                "E/G NodeB ID": 1001,
                "Cell ID": 2,
                "Site Lat": None,
                "Site Long": -75.2,
                "Azimuth": None,
            },
        ])
        normalized = records[["Latitude", "Longitude", "Azimuth"]].copy()

        resolved, metadata = ImportWizardDialog.resolve_cell_site_fields(
            normalized,
            records,
            cell_sites,
            {"Site ID": "Start_eNodeB", "Sector ID": "Start_Sector"},
            {
                "Site ID": "E/G NodeB ID",
                "Sector ID": "Cell ID",
                "Latitude": "Site Lat",
                "Longitude": "Site Long",
                "Azimuth": "Azimuth",
            },
            "cell_site_first_fallback_original",
        )

        self.assertEqual([45.1, 40.2, 40.3, 40.4], resolved["Latitude"].tolist())
        self.assertEqual([-75.1, -75.2, -70.3, -70.4], resolved["Longitude"].tolist())
        self.assertEqual([120, 20, 30, 40], resolved["Azimuth"].tolist())
        self.assertEqual(2, metadata["matched_rows"])
        self.assertEqual(1, metadata["unmatched_rows"])
        self.assertEqual(1, metadata["missing_key_rows"])
        self.assertEqual([3], metadata["unmatched_row_positions"])
        self.assertEqual([4], metadata["missing_key_row_positions"])
        self.assertEqual(
            {"Latitude": 1, "Longitude": 2, "Azimuth": 1},
            metadata["fields_from_cell_site_list"],
        )

        cell_site_only, _metadata = ImportWizardDialog.resolve_cell_site_fields(
            normalized,
            records,
            cell_sites,
            {"Site ID": "Start_eNodeB", "Sector ID": "Start_Sector"},
            {
                "Site ID": "E/G NodeB ID",
                "Sector ID": "Cell ID",
                "Latitude": "Site Lat",
                "Longitude": "Site Long",
                "Azimuth": "Azimuth",
            },
            "cell_site_only",
        )
        self.assertTrue(pd.isna(cell_site_only.loc[1, "Latitude"]))
        self.assertTrue(pd.isna(cell_site_only.loc[2, "Longitude"]))

        original_first, _metadata = ImportWizardDialog.resolve_cell_site_fields(
            normalized,
            records,
            cell_sites,
            {"Site ID": "Start_eNodeB", "Sector ID": "Start_Sector"},
            {
                "Site ID": "E/G NodeB ID",
                "Sector ID": "Cell ID",
                "Latitude": "Site Lat",
                "Longitude": "Site Long",
                "Azimuth": "Azimuth",
            },
            "original_first_fallback_cell_site",
        )
        self.assertEqual(records["Latitude"].tolist(), original_first["Latitude"].tolist())
        self.assertEqual(records["Longitude"].tolist(), original_first["Longitude"].tolist())
        self.assertEqual(records["Azimuth"].tolist(), original_first["Azimuth"].tolist())

    def test_import_can_resolve_tower_fields_from_separate_cell_site_list(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            records_path = root / "timing_advance.csv"
            cell_sites_path = root / "cell_sites.csv"
            pd.DataFrame([{
                "Start_DateTime": "2024-01-15 14:00:00",
                "Start_eNodeB": 1001,
                "Start_Sector": 2,
                "Latitude": 1.0,
                "Longitude": 2.0,
                "Azimuth": 3,
                "Start_Timing_Advance_Miles": 0.31,
            }]).to_csv(records_path, index=False)
            pd.DataFrame([{
                "E/G NodeB ID": 1001,
                "Cell ID": 2,
                "Azimuth": 135,
                "Site Lat": 43.15,
                "Site Long": -77.61,
            }]).to_csv(cell_sites_path, index=False)

            dialog = ImportWizardDialog()
            dialog.load_file(str(records_path), background=False)
            self.assertEqual(
                "Start_Timing_Advance_Miles",
                dialog.mapping_combos["Distance"].currentData(),
            )
            self.assertEqual(
                "Distance from Tower", dialog.record_type_combo.currentData()
            )
            dialog.use_cell_site_list_checkbox.setChecked(True)
            self.assertEqual(
                "cell_site_only", dialog.cell_site_policy_combo.currentData()
            )
            self.assertFalse(dialog.cell_site_policy_combo.isEnabled())
            for field in ("Latitude", "Longitude", "Azimuth"):
                self.assertFalse(dialog.mapping_combos[field].isEnabled())
                self.assertEqual(
                    "Not used",
                    dialog.cell_site_original_field_labels[field].text(),
                )
            self.assertEqual(
                {"Timestamp", "Distance"},
                set(dialog.current_mappings()),
            )
            dialog.load_cell_site_file(str(cell_sites_path), background=False)

            self.assertEqual(
                "Start_eNodeB",
                dialog.original_key_mapping_combos["Site ID"].currentData(),
            )
            self.assertEqual(
                "Start_Sector",
                dialog.original_key_mapping_combos["Sector ID"].currentData(),
            )
            self.assertEqual(
                "E/G NodeB ID",
                dialog.cell_site_mapping_combos["Site ID"].currentData(),
            )
            self.assertEqual(
                "Cell ID",
                dialog.cell_site_mapping_combos["Sector ID"].currentData(),
            )

            dialog.filter_enabled_checkbox.setChecked(True)
            with patch("import_wizard.QMessageBox.warning") as warning:
                dialog.preview_filter_results()
            warning.assert_not_called()

            dialog.accept_import()

            self.assertEqual(43.15, dialog.normalized_dataframe.loc[0, "Latitude"])
            self.assertEqual(-77.61, dialog.normalized_dataframe.loc[0, "Longitude"])
            self.assertEqual(135, dialog.normalized_dataframe.loc[0, "Azimuth"])
            self.assertEqual(0.31, dialog.normalized_dataframe.loc[0, "Distance"])
            self.assertEqual(1, dialog.selected_cell_site_metadata["matched_rows"])
            self.assertEqual([], dialog.selected_cell_site_metadata["unmatched_source_rows"])
            self.assertEqual([], dialog.selected_cell_site_metadata["missing_key_source_rows"])
            self.assertEqual(
                "cell_site_only",
                dialog.selected_cell_site_metadata["policy"],
            )
            self.assertEqual(
                ["Site ID", "Latitude", "Longitude", "CSL Source Rows"],
                list(dialog.reference_sites_dataframe.columns),
            )
            self.assertEqual(
                {
                    "Site ID": 1001.0,
                    "Latitude": 43.15,
                    "Longitude": -77.61,
                    "CSL Source Rows": (2,),
                },
                dialog.reference_sites_dataframe.iloc[0].to_dict(),
            )

            dialog.mapping_combos["Timestamp"].setCurrentIndex(0)
            with patch("import_wizard.QMessageBox.warning") as warning:
                dialog.preview_filter_results()
            warning.assert_called_once_with(
                dialog,
                "Incomplete Mapping",
                "Map the following fields before checking the filter:\n\n"
                "Timestamp (combined)",
            )

    def test_mapped_equivalent_cell_site_duplicates_are_collapsed(self):
        records = pd.DataFrame([{"Node": 1001, "Sector": 1}])
        normalized = pd.DataFrame(index=records.index)
        cell_sites = pd.DataFrame([
            {"Node": 1001, "Sector": 1, "Lat": 43.1, "Lon": -77.1, "Az": 120},
            {"Node": "1001", "Sector": "1", "Lat": "43.10", "Lon": "-77.10", "Az": "120.0"},
            {"Node": 2002, "Sector": 2, "Lat": 44.1, "Lon": -78.1, "Az": 20},
            {"Node": 2002, "Sector": 2, "Lat": 45.1, "Lon": -79.1, "Az": 30},
        ])

        resolved, metadata = ImportWizardDialog.resolve_cell_site_fields(
            normalized,
            records,
            cell_sites,
            {"Site ID": "Node", "Sector ID": "Sector"},
            {
                "Site ID": "Node",
                "Sector ID": "Sector",
                "Latitude": "Lat",
                "Longitude": "Lon",
                "Azimuth": "Az",
            },
            "cell_site_only",
        )

        self.assertEqual(43.1, resolved.loc[0, "Latitude"])
        self.assertEqual(-77.1, resolved.loc[0, "Longitude"])
        self.assertEqual(120, resolved.loc[0, "Azimuth"])
        self.assertEqual(1, metadata["duplicate_keys_collapsed"])
        self.assertEqual(1, metadata["duplicate_rows_collapsed"])
        self.assertEqual(1, metadata["unreferenced_conflicting_keys_ignored"])

    def test_reference_sites_omit_conflicting_and_incomplete_lookup_keys(self):
        cell_sites = pd.DataFrame([
            {"Node": 1001, "Sector": 1, "Lat": 43.1, "Lon": -77.1},
            {"Node": "1001", "Sector": "1", "Lat": "43.10", "Lon": "-77.10"},
            {"Node": 2002, "Sector": 2, "Lat": 44.1, "Lon": -78.1},
            {"Node": 2002, "Sector": 2, "Lat": 45.1, "Lon": -79.1},
            {"Node": 3003, "Sector": None, "Lat": 43.2, "Lon": -77.2},
            {"Node": 4004, "Sector": 4, "Lat": 43.3, "Lon": -77.3},
        ])

        references = ImportWizardDialog.build_reference_site_dataframe(
            cell_sites,
            {
                "Site ID": "Node",
                "Sector ID": "Sector",
                "Latitude": "Lat",
                "Longitude": "Lon",
            },
            cell_site_header_row=4,
        )

        self.assertEqual(2, len(references))
        self.assertEqual(
            {1001, 4004},
            {int(value) for value in references["Site ID"]},
        )
        rows_by_site = {
            int(row["Site ID"]): row["CSL Source Rows"]
            for _, row in references.iterrows()
        }
        self.assertEqual((5, 6), rows_by_site[1001])
        self.assertEqual((10,), rows_by_site[4004])

    def test_conflicting_used_cell_site_key_reports_key_rows_and_fields(self):
        records = pd.DataFrame([{"Node": 1001, "Sector": 1}])
        normalized = pd.DataFrame([{
            "Latitude": 40.1,
            "Longitude": -70.1,
            "Azimuth": 10,
        }])
        cell_sites = pd.DataFrame([
            {"Node": 1001, "Sector": 1, "Lat": 43.1, "Lon": -77.1},
            {"Node": 1001, "Sector": 1, "Lat": 44.1, "Lon": -78.1},
        ])

        with self.assertRaises(ValueError) as raised:
            ImportWizardDialog.resolve_cell_site_fields(
                normalized,
                records,
                cell_sites,
                {"Site ID": "Node", "Sector ID": "Sector"},
                {
                    "Site ID": "Node",
                    "Sector ID": "Sector",
                    "Latitude": "Lat",
                    "Longitude": "Lon",
                    "Azimuth": None,
                },
                "cell_site_only",
            )

        message = str(raised.exception)
        self.assertIn("conflicting tower data", message)
        self.assertIn("Site / Node ID 1001", message)
        self.assertIn("Sector / Cell ID 1", message)
        self.assertIn("cell site rows 2, 3", message)
        self.assertIn("Cell tower/site latitude: 43.1 vs 44.1", message)
        self.assertIn("Cell tower/site longitude: -77.1 vs -78.1", message)
        self.assertIn("Identical mapped duplicates are accepted automatically", message)

    def test_cell_site_excel_uses_its_own_sheet_and_header_row(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "cell_sites.xlsx"
            workbook = Workbook()
            first_sheet = workbook.active
            first_sheet.title = "Notes"
            first_sheet.append(["not the cell site list"])
            cell_site_sheet = workbook.create_sheet("Cell Sites")
            cell_site_sheet.append(["Report title"])
            cell_site_sheet.append([
                "E/G NodeB ID", "Cell ID", "Site Lat", "Site Long", "Azimuth"
            ])
            cell_site_sheet.append([1001, 2, 43.15, -77.61, 135])
            workbook.save(path)

            dialog = ImportWizardDialog()
            dialog.record_type_combo.setCurrentIndex(
                dialog.record_type_combo.findData("Distance from Tower")
            )
            dialog.use_cell_site_list_checkbox.setChecked(True)
            dialog._cell_site_header_reload_timer.setInterval(0)
            dialog.load_cell_site_file(str(path), background=False)
            dialog.cell_site_sheet_combo.setCurrentText("Cell Sites")
            dialog.cell_site_header_row_spinbox.setValue(2)
            self.app.processEvents()

            self.assertEqual(
                ["E/G NodeB ID", "Cell ID", "Site Lat", "Site Long", "Azimuth"],
                list(dialog.cell_site_dataframe.columns),
            )
            self.assertEqual(
                "E/G NodeB ID",
                dialog.cell_site_mapping_combos["Site ID"].currentData(),
            )

            dialog.record_type_combo.setCurrentIndex(
                dialog.record_type_combo.findData("Location Point")
            )
            self.assertFalse(dialog.use_cell_site_list_checkbox.isEnabled())
            self.assertFalse(dialog.cell_site_list_enabled())


if __name__ == "__main__":
    unittest.main()
